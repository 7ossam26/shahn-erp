"""Independent output inspection; run with the bundled Python and Poppler paths."""
from pathlib import Path
import json, re, subprocess, sys
from pypdf import PdfReader
from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageDraw

root = Path('docs/verification/P23')
poppler = Path(sys.argv[1])
renders = root / 'pdf-pages'
renders.mkdir(exist_ok=True)
assert renders.resolve().parent == root.resolve()
for old in renders.glob('*.png'):
    old.unlink()
checks = []
for pdf in sorted((root / 'reports').glob('*.pdf')):
    workbook = pdf.with_suffix('.xlsx')
    book = load_workbook(workbook, data_only=False)
    source = book['المصادر']
    snapshot_cell = str(book['التقرير'].cell(5,1).value)
    snapshot_id = re.search(r'snapshot ([a-f0-9-]+)', snapshot_cell).group(1)
    reader = PdfReader(pdf)
    text = '\n'.join(p.extract_text() for p in reader.pages)
    assert snapshot_id in text.replace('\n',''), (pdf, 'snapshot missing')
    assert all(c.data_type != 'f' for s in book for row in s for c in row), 'formula cell'
    prefix = renders / pdf.stem
    subprocess.run([str(poppler / 'pdftoppm.exe'), '-png', '-scale-to', '1800', str(pdf), str(prefix)], check=True, capture_output=True)
    checks.append({'path':str(pdf),'pages':len(reader.pages),'rows':source.max_row-1,'snapshotId':snapshot_id,'workbook':str(workbook),'noExcelFormulaCells':True})
for pdf, workbook in [('shipments-27.pdf','same-snapshot-27.xlsx'), ('expenses-200.pdf','same-snapshot-expenses.xlsx'), ('returns-2-1.pdf','returns-2-1.xlsx'), ('stock-transit.pdf','stock-transit.xlsx')]:
    path = root / pdf
    book = load_workbook(root / workbook, data_only=False)
    reader = PdfReader(path)
    text = '\n'.join(p.extract_text() for p in reader.pages)
    snapshot_id = re.search(r'snapshot ([a-f0-9-]+)', str(book['التقرير'].cell(5,1).value)).group(1)
    assert snapshot_id in text.replace('\n',''), 'snapshot mismatch'
    if pdf == 'shipments-27.pdf':
        refs = [str(book['التقرير'].cell(i,1).value) for i in range(8,35)]
        assert len(set(refs)) == 27
        assert all(re.search(r'(?<!\d)'+re.escape(x)+r'(?!\d)', text) for x in refs), 'missing PDF reference'
    if pdf == 'expenses-200.pdf':
        assert '200.00' in text
        assert book['التقرير'].cell(8,6).value == '200.00'
    subprocess.run([str(poppler / 'pdftoppm.exe'), '-png', '-scale-to', '1800', str(path), str(renders / path.stem)], check=True, capture_output=True)
    checks.append({'path':str(path),'pages':len(reader.pages),'rows':book['المصادر'].max_row-1,'snapshotId':snapshot_id,'workbook':workbook})
images = list(sorted(p for p in renders.glob('*.png') if not p.name.startswith('contact-')))
# Contact sheets index every rendered page; individual PNGs remain available for full-size review.
for start in range(0,len(images),6):
    canvas = Image.new('RGB',(1800,1800),'white')
    draw = ImageDraw.Draw(canvas)
    for i,path in enumerate(images[start:start+6]):
        im = Image.open(path).convert('RGB'); im.thumbnail((880,550))
        x=(i%2)*900; y=(i//2)*600
        canvas.paste(im,(x,y+35)); draw.text((x+10,y+10),path.name,fill='black')
    canvas.save(renders / f'contact-{start//6+1}.png')
(root/'artifact-inspection.json').write_text(json.dumps({'pdfs':checks,'renderedPages':len(images),'checks':'parsed snapshot identity, every selected report XLSX cells, PDF27 references and expense200; visual QA separately recorded'},indent=2),encoding='utf-8')
print(json.dumps({'pdfs':len(checks),'renderedPages':len(images),'output':str(root/'artifact-inspection.json')}))
