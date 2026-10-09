"""Independent XLSX/PDF identity, exact money and page rendering checks for P24."""
from pathlib import Path
import json, re, subprocess, sys
from pypdf import PdfReader
from openpyxl import load_workbook
from PIL import Image, ImageDraw

root = Path('docs/verification/P24')
poppler = Path(sys.argv[1])
renders = root / 'pdf-pages'
renders.mkdir(exist_ok=True)
checks = []
for pdf in sorted(root.glob('*.pdf')):
    workbook = pdf.with_suffix('.xlsx')
    assert workbook.exists(), f'Workbook missing: {pdf}'
    book = load_workbook(workbook, data_only=False)
    report, source, context = book['التقرير'], book['المصادر'], book['سياق وأرصدة']
    snapshot_id = re.search(r'snapshot ([a-f0-9-]+)', str(report.cell(5, 1).value)).group(1)
    reader = PdfReader(pdf)
    text = '\n'.join(p.extract_text() for p in reader.pages)
    assert snapshot_id in text.replace('\n',''), f'Snapshot mismatch: {pdf}'
    assert all(c.data_type != 'f' for s in book for row in s for c in row), f'Formula cell: {workbook}'
    assert report.sheet_view.rightToLeft, f'RTL missing: {workbook}'
    raw_context = {str(context.cell(i,1).value):str(context.cell(i,2).value) for i in range(1,context.max_row+1)}
    profit = json.loads(raw_context['profit'])
    amounts = [json.loads(source.cell(i,6).value)['amountMinor'] for i in range(2,source.max_row+1)]
    total = sum(int(n) for n in amounts)
    assert str(total) == profit['profitMinor'], f'Source sum mismatch: {workbook}'
    assert sum(int(b['profitMinor']) for b in profit['branches']) == total, f'Branch sum mismatch: {workbook}'
    assert sum(int(c['amountMinor']) for c in profit['categories']) == total, f'Category sum mismatch: {workbook}'
    def display(n):
        a=abs(int(n)); return f'{a//100}.{a%100:02d}'
    for n in amounts + [str(total)]:
        assert display(n) in text, f'Exact amount missing from PDF: {pdf} amount {n}'
    assert all(isinstance(source.cell(i,1).value,str) for i in range(2,source.max_row+1)), 'Source IDs not text'
    assert all(source.cell(i,7).value and json.loads(source.cell(i,7).value)['sourceId'] for i in range(2,source.max_row+1)), 'Missing economic source provenance'
    subprocess.run([str(poppler/'pdftoppm.exe'),'-png','-scale-to','1800',str(pdf),str(renders/pdf.stem)],check=True,capture_output=True)
    checks.append({'pdf':str(pdf),'workbook':str(workbook),'snapshotId':snapshot_id,'pages':len(reader.pages),'rows':source.max_row-1,'profitMinor':str(total),'rtl':True,'formulaCells':0,'sourceCategoryBranchTotalsAgree':True,'exactPDFValuesPresent':True})
images = sorted(p for p in renders.glob('*.png') if not p.name.startswith('contact-'))
for start in range(0,len(images),4):
    canvas=Image.new('RGB',(1800,1400),'white'); draw=ImageDraw.Draw(canvas)
    for i,path in enumerate(images[start:start+4]):
        im=Image.open(path).convert('RGB'); im.thumbnail((880,640))
        x=(i%2)*900; y=(i//2)*700
        canvas.paste(im,(x,y+35)); draw.text((x+10,y+10),path.name,fill='black')
    canvas.save(renders/f'contact-{start//4+1}.png')
result={'pdfs':checks,'renderedPages':len(images),'visualInspection':'Record separately after viewing every rendered page.'}
(root/'artifact-inspection.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
